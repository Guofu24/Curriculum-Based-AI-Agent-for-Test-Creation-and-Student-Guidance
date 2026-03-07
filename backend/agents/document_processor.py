"""
Document Processor Agent

Responsibilities:
- Parse uploaded documents (PDF, DOCX, PPTX)
- Chunk text with overlap for embedding
- Extract chapter structure and metadata
- Store embeddings in Pinecone (cloud) and raw text in PostgreSQL
"""
import os
import json
import hashlib
from pathlib import Path

from langchain_core.documents import Document

from config import settings


class DocumentProcessorAgent:
    """Processes textbook documents: parse, chunk, embed, store in Pinecone."""

    CHUNK_SIZE = 1000  # characters per chunk
    CHUNK_OVERLAP = 200  # overlap between chunks

    def __init__(self, vector_store, db_session=None):
        self.vector_store = vector_store
        self.db_session = db_session  # SQLAlchemy session for storing chunk text

    async def process_document(self, file_path: str, textbook_id: str) -> dict:
        """
        Full document processing pipeline.
        Returns metadata about the processed textbook.
        """
        # 1. Parse the document
        raw_documents = await self._parse_document(file_path)

        # 2. Extract chapter structure
        chapters = self._extract_chapters(raw_documents)

        # 3. Chunk the documents
        chunks = self._chunk_documents(raw_documents, textbook_id)

        # 4. Store in vector DB with embeddings
        chunk_ids = await self._store_chunks(chunks, textbook_id)

        return {
            "textbook_id": textbook_id,
            "title": Path(file_path).stem,
            "total_pages": len(raw_documents),
            "total_chunks": len(chunks),
            "chapters": chapters,
            "chunk_ids": chunk_ids,
        }

    async def _parse_document(self, file_path: str) -> list[Document]:
        """Parse document based on file type."""
        ext = Path(file_path).suffix.lower()

        if ext == ".pdf":
            return await self._parse_pdf(file_path)
        elif ext == ".docx":
            return await self._parse_docx(file_path)
        elif ext in (".pptx", ".ppt"):
            return await self._parse_pptx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")

    async def _parse_pdf(self, file_path: str) -> list[Document]:
        from pypdf import PdfReader

        reader = PdfReader(file_path)
        documents = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                documents.append(Document(
                    page_content=text,
                    metadata={"page": i + 1, "source": file_path}
                ))
        return documents

    async def _parse_docx(self, file_path: str) -> list[Document]:
        from docx import Document as DocxDoc

        doc = DocxDoc(file_path)
        documents = []
        current_text = []
        page_num = 1

        for para in doc.paragraphs:
            if para.text.strip():
                current_text.append(para.text)
            # Approximate page breaks every ~3000 chars
            combined = "\n".join(current_text)
            if len(combined) > 3000:
                documents.append(Document(
                    page_content=combined,
                    metadata={"page": page_num, "source": file_path}
                ))
                current_text = []
                page_num += 1

        if current_text:
            documents.append(Document(
                page_content="\n".join(current_text),
                metadata={"page": page_num, "source": file_path}
            ))
        return documents

    async def _parse_pptx(self, file_path: str) -> list[Document]:
        from pptx import Presentation

        prs = Presentation(file_path)
        documents = []
        for i, slide in enumerate(prs.slides):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text)
            if texts:
                documents.append(Document(
                    page_content="\n".join(texts),
                    metadata={"page": i + 1, "source": file_path, "slide": i + 1}
                ))
        return documents

    def _extract_chapters(self, documents: list[Document]) -> list[dict]:
        """
        Heuristic chapter extraction from document structure.
        Looks for patterns like "Chapter X", "CHAPTER X", "Chương X".
        """
        import re
        chapters = []
        chapter_pattern = re.compile(
            r"(?:chapter|chương|ch\.?)\s*(\d+)[:\s.\-]*(.+)",
            re.IGNORECASE
        )

        for doc in documents:
            for line in doc.page_content.split("\n"):
                match = chapter_pattern.match(line.strip())
                if match:
                    chapter_num = int(match.group(1))
                    chapter_title = match.group(2).strip()
                    if not any(c["chapter_number"] == chapter_num for c in chapters):
                        chapters.append({
                            "chapter_number": chapter_num,
                            "title": chapter_title,
                            "start_page": doc.metadata.get("page", 0),
                        })

        chapters.sort(key=lambda c: c["chapter_number"])

        # Set end_page for each chapter
        for i, ch in enumerate(chapters):
            if i + 1 < len(chapters):
                ch["end_page"] = chapters[i + 1]["start_page"] - 1
            else:
                ch["end_page"] = documents[-1].metadata.get("page", 0) if documents else 0

        return chapters

    def _chunk_documents(
        self, documents: list[Document], textbook_id: str
    ) -> list[Document]:
        """Split documents into overlapping chunks with metadata."""
        chunks = []
        chunk_idx = 0

        for doc in documents:
            text = doc.page_content
            start = 0
            while start < len(text):
                end = start + self.CHUNK_SIZE
                chunk_text = text[start:end]

                if chunk_text.strip():
                    chunk_id = hashlib.md5(
                        f"{textbook_id}:{chunk_idx}".encode()
                    ).hexdigest()

                    chunks.append(Document(
                        page_content=chunk_text,
                        metadata={
                            **doc.metadata,
                            "textbook_id": textbook_id,
                            "chunk_id": chunk_id,
                            "chunk_index": chunk_idx,
                        }
                    ))
                    chunk_idx += 1

                start += self.CHUNK_SIZE - self.CHUNK_OVERLAP

        return chunks

    async def _store_chunks(
        self, chunks: list[Document], textbook_id: str
    ) -> list[str]:
        """Store chunks in Pinecone (vectors) and PostgreSQL (raw text for BM25)."""
        texts = [c.page_content for c in chunks]
        metadatas = [c.metadata for c in chunks]
        ids = [c.metadata["chunk_id"] for c in chunks]

        # 1. Store embeddings in Pinecone via LangChain vector store (async to avoid blocking event loop)
        await self.vector_store.aadd_texts(
            texts=texts,
            metadatas=metadatas,
            ids=ids,
        )

        # 2. Store raw text in PostgreSQL for BM25 keyword search
        if self.db_session is not None:
            import json
            from models.textbook import TextbookChunk

            for i, chunk in enumerate(chunks):
                db_chunk = TextbookChunk(
                    textbook_id=textbook_id,
                    chunk_id=ids[i],
                    chunk_index=i,
                    content=texts[i],
                    page=chunk.metadata.get("page"),
                    metadata_json=json.dumps(metadatas[i], ensure_ascii=False),
                )
                self.db_session.add(db_chunk)

        return ids

    def _assign_chapter_to_chunk(
        self, page: int, chapters: list[dict]
    ) -> int:
        """Determine which chapter a page belongs to."""
        for ch in reversed(chapters):
            if page >= ch.get("start_page", 0):
                return ch["chapter_number"]
        return 1  # default to chapter 1
